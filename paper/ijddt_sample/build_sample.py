"""Rebuild the reference-styled sample and its six-page manuscript copy.

Run with the Codex bundled Python. PDF export and visual QA are separate steps.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

OUT = Path(__file__).resolve().parent
doc = Document()
markdown = []

def setup(section, cols=1):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(.65)
    section.footer_distance = Cm(.7)
    c = section._sectPr.find(qn('w:cols'))
    if c is None:
        c = OxmlElement('w:cols'); section._sectPr.append(c)
    c.set(qn('w:num'), str(cols)); c.set(qn('w:space'), '360')
    section.different_first_page_header_footer = False

setup(doc.sections[0])
for name, size in [('Normal',10.5), ('Title',14), ('Heading 1',11), ('Heading 2',10.5), ('Caption',9.5)]:
    s = doc.styles[name]
    s.font.name = 'Times New Roman'; s.font.size = Pt(size); s.font.color.rgb = RGBColor(0,0,0)
    fonts = s.element.get_or_add_rPr().find(qn('w:rFonts'))
    for attr in list(fonts.attrib):
        if attr.endswith('Theme'): del fonts.attrib[attr]
    for attr in ['ascii','hAnsi','eastAsia','cs']: fonts.set(qn('w:'+attr),'Times New Roman')
    sp = s.element.find(qn('w:pPr'))
    if sp is not None:
        for border in list(sp.findall(qn('w:pBdr'))): sp.remove(border)
    s.paragraph_format.space_after = Pt(5)
    s.paragraph_format.line_spacing = Pt(12.5)
    s.paragraph_format.widow_control = True
    if name in ['Heading 1','Heading 2']:
        s.font.bold = True
        s.paragraph_format.space_before = Pt(7)
        s.paragraph_format.keep_with_next = True
    if name == 'Title':
        s.font.bold = True; s.paragraph_format.line_spacing = Pt(16)
        s.paragraph_format.space_after = Pt(9)
doc.styles['Caption'].font.italic = False
doc.styles['Caption'].font.bold = False
doc.core_properties.title = 'Evidence linked Graph RAG for clinical claim review'
doc.core_properties.subject = 'IJDDT reference-styled sample manuscript with synthetic evaluation results'
doc.core_properties.author = 'Sample authors - replace before submission'

h = doc.sections[0].header.paragraphs[0]
h.alignment = WD_ALIGN_PARAGRAPH.CENTER
f=doc.sections[0].footer.paragraphs[0]
f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); f._p.append(field)

def p(text, style=None, align=None, size=None, italic=False, bold=False):
    para=doc.add_paragraph(style=style)
    para.alignment=align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    # Citation markers are authored as superscript numeric Vancouver citations.
    for part in re.split(r'(\[\d+(?:,\d+)*\])',text):
        if not part: continue
        r=para.add_run(part[1:-1] if re.fullmatch(r'\[\d+(?:,\d+)*\]',part) else part)
        if re.fullmatch(r'\[\d+(?:,\d+)*\]',part): r.font.superscript=True
        if size: r.font.size=Pt(size)
        if italic: r.italic=True
        if bold: r.bold=True
        if style == 'Title':
            fonts=r._r.get_or_add_rPr().find(qn('w:rFonts'))
            if fonts is None: fonts=OxmlElement('w:rFonts');r._r.get_or_add_rPr().append(fonts)
            for attr in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+attr),'Times New Roman')
    markdown.append(text+'\n')
    return para

def head(text, level=1):
    para=p(text, 'Heading '+str(level), WD_ALIGN_PARAGRAPH.LEFT)
    markdown[-1]='#'*(level+1)+' '+text+'\n'
    return para

def section(cols=1,newpage=False):
    s=doc.add_section(WD_SECTION_START.NEW_PAGE if newpage else WD_SECTION_START.CONTINUOUS)
    setup(s,cols)
    # Section break paragraphs should not consume a full body line.
    prev=doc.paragraphs[-1]
    prev.paragraph_format.space_after=Pt(0); prev.paragraph_format.space_before=Pt(0)
    prev.paragraph_format.line_spacing=Pt(1)
    for r in prev.runs:r.font.size=Pt(1)
    return s

def colbreak():
    para=doc.add_paragraph()
    para.paragraph_format.space_after=Pt(0); para.paragraph_format.line_spacing=Pt(1)
    para.add_run().add_break(WD_BREAK.COLUMN)

def placeholder(label, description, height):
    para=p(label+'\n'+description,align=WD_ALIGN_PARAGRAPH.CENTER,size=10,italic=True)
    # A fixed paragraph slot leaves usable space without a floating drawing.
    para.paragraph_format.space_before=Pt(height*72/2-12)
    para.paragraph_format.space_after=Pt(height*72/2-12)
    para.paragraph_format.keep_with_next=True
    markdown[-1]='[Reserved figure space: '+label+'; '+description+']\n'

def table(headers, rows, widths):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Cm(w)
    for i,text in enumerate(headers):t.rows[0].cells[i].text=text
    for row in rows:
        cells=t.add_row().cells
        for i,text in enumerate(row):cells[i].text=str(text)
    for ri,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr()
        no=OxmlElement('w:cantSplit');trPr.append(no)
        if ri==0:
            rep=OxmlElement('w:tblHeader');trPr.append(rep)
        for ci,c in enumerate(row.cells):
            c.width=Cm(widths[ci]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tcPr.append(borders)
            margins=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'70');e.set(qn('w:type'),'dxa');margins.append(e)
            tcPr.append(margins)
            if ri==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F0F0F0');tcPr.append(shade)
            for para in c.paragraphs:
                para.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                para.paragraph_format.space_after=Pt(0);para.paragraph_format.line_spacing=Pt(11.5)
                for r in para.runs:r.font.size=Pt(9.5);r.bold=(ri==0)
    markdown.append('| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |\n'+''.join('| '+' | '.join(map(str,row))+' |\n' for row in rows))
    doc.add_paragraph().paragraph_format.space_after=Pt(0)
    return t

# PAGE 1: full-width front matter, then two-column Introduction.
p('RESEARCH PAPER',align=WD_ALIGN_PARAGRAPH.LEFT,bold=True,size=10.5)
p('EVIDENCE LINKED GRAPH RAG FOR CLINICAL CLAIM REVIEW\nA SYNTHETIC FEASIBILITY STUDY',style='Title',align=WD_ALIGN_PARAGRAPH.CENTER)
p('Aarav Sharma¹*, Meera Rao¹, Priya Nair², Rohan Verma¹ and Kavya Iyer²',align=WD_ALIGN_PARAGRAPH.CENTER,bold=True,size=12)
p('¹ Department of Computer Science and Engineering, Example Institute of Technology, Hyderabad, India\n² Department of Pharmacy Practice, Example College of Pharmacy, Hyderabad, India',align=WD_ALIGN_PARAGRAPH.CENTER,size=10.5,italic=True)
p('*Corresponding author: author@example.com',align=WD_ALIGN_PARAGRAPH.CENTER,size=10)
head('ABSTRACT')
p('Retrieval-augmented clinical decision support can produce plausible findings with citations that do not establish patient-specific support. We describe a research prototype that separates exploratory evidence from passages eligible for citation. Patient facts and missing information are extracted from a free-text case. Biomedical graph relations and textbook vectors inform claim generation, while selected NICE hypertension recommendations and one DailyMed amlodipine label form a versioned authoritative corpus. A separate verifier checks cited passages against local source snapshots, document metadata and SHA-256 hashes. General claims remain unverified pending clinical review; two narrow medication-safety rules can produce source-backed findings. We evaluated the backend in LLM-only, text-RAG and graph-plus-text modes. A historical six-case fixture set completed 18 analyses. A separate 12-case, diagnosis-withheld pilot completed 36 analyses, with exact provisional-label matches at rank 1 of 10/12, 9/12 and 10/12, respectively. Citation identity matched 30/30 and 26/26 emitted uses in the pilot retrieval arms. All pilot claims remained unverified, and the graph returned no disease candidates. These results demonstrate an inspectable evidence workflow but do not establish diagnostic superiority, treatment safety or clinical utility. Broader source coverage, controlled retrieval evaluation and independent clinical adjudication are required.')
p('Keywords: Clinical decision support; retrieval-augmented generation; biomedical knowledge graph; source provenance; medication safety; synthetic evaluation.',size=10,italic=True)
section(2)
head('1. INTRODUCTION')
p('A clinical decision-support report is useful only when the reviewer can connect its findings to the patient facts and evidence on which they depend. A generated explanation may sound reasonable even when an allergy is unknown, a guideline addresses a different population, or a cited passage concerns treatment rather than diagnosis. These errors can persist despite successful document retrieval.')
p('Retrieval-augmented generation combines a language model with external information supplied during inference.[1] In a clinical setting, however, finding a relevant-looking passage is only one part of the task. The system must also preserve the source version, distinguish missing information from negative findings, and make unresolved applicability checks visible.')
colbreak()
p('We therefore examine a pipeline for reviewing clinical claims rather than issuing autonomous prescribing decisions. The design retains a broad exploratory channel for candidate generation and a narrow, pinned-source channel for citations. A verifier operates after synthesis so that a model cannot establish source validity merely by attaching an identifier to its own output.')
p('The study asks whether this separation can be implemented and audited in a reproducible research prototype. Its contribution is a claim-level provenance and abstention workflow, accompanied by synthetic engineering evaluations. The evaluation is intentionally limited: neither valid citation identity nor a match to an author-written diagnosis label is treated as proof of clinical correctness.')

# PAGE 2: background and first half of methods.
section(2,True)
head('2. RELATED WORK / LITERATURE REVIEW')
p('Retrieval-augmented generation was introduced to combine parametric generation with a retrievable external corpus.[1] It provides a way to inspect supporting material, but it does not by itself guarantee that the generated statement follows from that material. Clinical systems need to account for patient eligibility, document scope and jurisdiction in addition to similarity between query and passage.')
p('Biomedical knowledge graphs offer a complementary representation. PrimeKG integrates heterogeneous biomedical entities and relations that can support disease and drug association searches.[2] Graph relations can help identify candidates for investigation, while text can provide the wording and conditions of a recommendation. A graph association is not equivalent to an applicable treatment recommendation, and imported graph coverage determines which paths are available.')
p('The present system combines graph and dense-text retrieval but keeps their outputs separate from its citable source records. It does not implement a new embedding model, a learned reranker or a general clinical entailment classifier. Its emphasis is on preserving source identity and making uncertainty explicit at the level of individual claims. Novelty relative to comparable evidence-grounding systems remains a question for a fuller literature comparison.')
head('3. METHODOLOGY / PROPOSED METHOD')
head('3.1 Task definition and system architecture',2)
p('The input is a free-text synthetic clinical case. The output is a structured report containing a patient summary, differential candidates, typed claims, medication warnings, exact source links and unresolved checks. A browser interface sends the case to a FastAPI analysis endpoint. Neo4j provides graph retrieval and Qdrant provides separate textbook, guideline and drug-label collections. Figure 1 presents the main workflow and its citation boundary.')
p('The runtime proceeds through context extraction, retrieval, evidence ranking, claim synthesis, independent verification and report assembly. The normal application can use local DrugCentral data and optional web discovery as additional leads. These channels are disabled in the reported ablations. The architecture should therefore distinguish the complete application from the narrower evaluation configuration.')
colbreak()
head('3.2 Patient context and entity linking',2)
p('A clarifier extracts symptoms, known conditions, medicines, allergy statements and demographic information. A separate fact-extraction pass anchors important values to the original wording. Local HPO, MONDO and drug-name resources support exact and fuzzy entity linking. The patient context retains unknown or conflicting states instead of treating omitted information as a reassuring negative.')
p('Anchoring reduces one source of error but cannot make extraction infallible. In the historical fixtures, an unknown-allergy statement was misread and the expected missingness warning was not produced. Downstream verification cannot compensate for a patient fact that was extracted incorrectly. The report consequently requires review of the patient summary as well as its citations.')
head('3.3 Exploratory graph and textbook retrieval',2)
p('Graph retrieval queries a locally imported biomedical graph for disease candidates and related signals. Textbook retrieval uses the sentence-transformer model pritamdeka/S-PubMedBert-MS-MARCO.[3] The ingestion script creates 768-dimensional vectors with cosine distance. It reads First Aid Step 1 from the MedQA textbook directory, truncates the input to 150,000 characters, and splits it into 150-word chunks with a 20-word overlap. This describes the supplied ingestion recipe; it is not a complete-corpus coverage claim.')
p('One dense textbook query combines the case or expanded query with extracted symptom text. The retriever requests five nearest neighbours and retains hits with a score greater than 0.4. This hard-coded threshold was not calibrated against a relevance-labelled dataset. No BM25 retrieval, reciprocal-rank fusion or cross-encoder reranking is implemented.')
p('The evidence ranker multiplies graph candidate scores by 1.5 and adds 0.1 to a textbook score for each normalized symptom name found as a case-insensitive substring in the passage. These are hand-set ordering rules. Their output is not interpreted as a diagnostic probability. Graph edges, textbook passages and discovery snippets remain exploratory signals and cannot independently establish supported clinical claims.')

# PAGE 3: figure slot and evidence boundary details.
section(1,True)
placeholder('INSERT FIGURE 1 HERE', 'System architecture showing parallel retrieval and the separate pinned-source verification path',2.35)
p('Figure 1. Proposed architecture for evidence-linked clinical claim review. Graph and textbook signals inform generation; pinned guideline and label passages also enter independent source verification. A dashed retry loop is optional in the application and disabled in the reported evaluations.',style='Caption',align=WD_ALIGN_PARAGRAPH.CENTER)
section(2)
head('3.4 Versioned authoritative retrieval',2)
p('The authoritative channel contains selected recommendations from NICE NG136 for adult hypertension in England[4] and one United States DailyMed amlodipine label.[5] The local snapshots yield 17 guideline passages and 9 label passages. The manifest records document identity, version, date, jurisdiction, publisher URL, local snapshot path and SHA-256 hash. These documents are not interchangeable worldwide standards.')
p('Guideline and label passages are indexed into dedicated Qdrant collections with the same sentence encoder as the textbook channel. At retrieval, the source manifest and local snapshots are checked before passage identifiers are accepted. Five hits are requested from each applicable source collection. Unlike textbook retrieval, this path does not apply the 0.4 score cutoff.')
p('A missing or changed snapshot prevents the source from being trusted. Retrieval returns limitations instead of silently replacing the pinned version. Source identifiers incorporate document identity, version, section, span and passage hash. A model proposal must cite an identifier that resolves to an included authoritative passage; invented or exploratory evidence identifiers are rejected.')
head('3.5 Claim generation',2)
p('The fusion stage receives the patient context, ranked candidate evidence and authoritative passages. It proposes differential entries and separate diagnosis, investigation and conditional-treatment claims. Each claim includes its intended target, conditions, missing checks, evidence identifiers and verification status. New model claims begin as unverified. Structured schemas check reference integrity when the final report is assembled.')
colbreak()
head('3.6 Independent verification and safety rules',2)
p('ClaimVerifier reconstructs the pinned passage set from the local snapshots and compares retrieved records with the pinned text and metadata. It checks every proposed citation and applies limited source-type, topic and population restrictions. A treatment recommendation cannot establish a diagnosis, and selected adult-hypertension treatment sections do not establish eligibility for a pregnant patient.')
p('Passing a source-identity check does not establish semantic support or patient-specific treatment suitability. The verifier therefore leaves general claims unverified. A separate medication-safety path applies two narrow rules concerning stated amlodipine sensitivity and amlodipine with simvastatin. A source-backed finding requires explicit patient facts and the corresponding pinned label passage. These rules do not constitute a comprehensive interaction checker.')
p('Unknown allergies, an incomplete medication list and other missing prerequisites remain visible as unresolved checks. The report includes exact passage links and verification failures so that a reviewer can examine both the proposal and the reason it was not established. An unavailable source produces a gap rather than permission to proceed.')
head('3.7 Adaptive retrieval and report generation',2)
p('In the application, an uncalibrated confidence heuristic may trigger up to two additional attempts. Expansion appends names from up to two linking candidates per symptom while preserving the original case as the fact-extraction anchor. It does not implement exhaustive ontology-synonym expansion. The report groups differential candidates, claims, medication findings and source passages, with conditions and missing information retained. The interface hides the legacy confidence percentage. Evaluation runs used zero retries so that comparisons remained bounded.')

# PAGE 4: experimental configuration and evaluation protocol.
section(1,True)
head('4. EXPERIMENTAL SETUP / DATASET')
p('Table 1. Retrieval and execution settings used by the inspected backend.',style='Caption',align=WD_ALIGN_PARAGRAPH.LEFT)
table(['Component','Current configuration','Interpretation'],[
    ['Textbook collection','medical_textbooks','One capped textbook ingestion recipe'],
    ['Encoder','S-PubMedBert-MS-MARCO','768 dimensions; cosine distance'],
    ['Textbook search','Top 5; score > 0.4','Heuristic cutoff; relevance not calibrated'],
    ['Pinned source collections','clinical_guidelines; clinical_drug_labels','NICE hypertension and one amlodipine label'],
    ['Evaluation modes','llm_only; text_rag; graph_text','Shared analysis path; retrieval substituted'],
    ['Evaluation settings','Temperature 0; retries 0','Web discovery and DrugCentral disabled'],
], [4.1,6.8,6.5])
section(2)
head('4.1 Synthetic datasets',2)
p('The historical dataset contains six hand-authored engineering fixtures. Four concern known hypertension and test amlodipine allergy, amlodipine with simvastatin, unknown allergy status and absence of an identified warning. A pregnancy case and a known gastro-oesophageal reflux disease case test boundaries outside the pinned adult-hypertension treatment domain. The reference diagnosis is explicitly written in every historical case, so diagnosis matching measures recognition rather than independent diagnostic inference.')
p('A separate pilot contains 12 author-built presentations with diagnosis names withheld: three each for hypertension, chronic asthma, type 2 diabetes and gastro-oesophageal reflux disease or dyspepsia. The labels are provisional author-intended diagnoses, not a clinician reference standard. These cases remain outside the indexed source corpus. The authoritative application corpus was not expanded to cover the additional pilot conditions.')
head('4.2 Retrieval ablations',2)
p('LLM-only removes graph, textbook, authoritative-passage, web and DrugCentral retrieval. Text RAG supplies textbook vectors and pinned guideline or label passages. Graph plus text adds graph disease candidates to the same text retrieval. Context extraction runs once per case and is reused across modes; the order of modes rotates by case. Fusion, verification, safety checks and report assembly are shared.')
p('The six-case run completed on 7 October 2026. The diagnosis-withheld pilot completed on 12 October 2026. Both used the backend analysis route through the evaluation runner. They are separate datasets with different diagnostic tasks and must not be pooled into a single accuracy estimate.')
colbreak()
head('4.3 Outcomes and reproducibility',2)
p('The diagnostic endpoint is strict normalized matching of an emitted diagnosis to an author-listed alias at rank 1 and within ranks 1-3. Citation validity measures whether emitted passage identifiers resolve to records that match the pinned corpus. Narrow warning counters are available in the historical fixtures. Claim abstention records the fraction of claims left unverified. Failures, completion counts and latency are recorded separately.')
p('Captured reports can be rescored without new model calls. The run manifest records dataset and source hashes, prompt hashes, model-provider selections, Git state and execution settings. Each case result retains the report, retrieval outputs, provider trace and elapsed time. Private mappings for blinded review are kept separate from reviewer-facing claim packets.')
p('Configured providers are attempted through a fallback mechanism. Actual synthesis-provider selections differed across modes for two pilot cases; an OpenRouter free route also prevents recovery of an exact underlying model version. A retrieval comparison is therefore not fully controlled merely because it uses the same configured fallback order.')
head('4.4 Unavailable measurements',2)
p('Neither dataset includes an independently labelled set of relevant passages. Retrieval Hit@k, Recall@k and mean reciprocal rank are consequently unavailable. The 12-case pilot includes no positive warning cases, so it cannot estimate warning sensitivity. Clinical claim support and unsafe-treatment rates require independent adjudication; they are not inferred from model status or citation counts.')
p('The local graph audit for the pilot found no phenotype-type relations needed by the current disease query, and no disease candidates were returned. The graph arm tests the configured request path under this coverage limitation. It does not establish the benefit of a functioning, sufficiently populated graph retriever.')

# PAGE 5: measured results and a report screenshot slot.
section(1,True)
head('5. RESULTS AND DISCUSSION')
p('Table 2. Descriptive outcomes from the two separate synthetic evaluations. Counts are label matches or source-identity checks, not clinical accuracy or treatment-safety estimates.',style='Caption',align=WD_ALIGN_PARAGRAPH.LEFT)
table(['Outcome','LLM only','Text RAG','Graph + text'],[
    ['Historical fixtures completed','6/6','6/6','6/6'],
    ['Historical rank-1 label match','1/6','4/6','5/6'],
    ['Historical rank-3 label match','1/6','4/6','6/6'],
    ['Pilot analyses completed','12/12','12/12','12/12'],
    ['Pilot rank-1 and rank-3 matches','10/12','9/12','10/12'],
    ['Pilot citation identity matched','No IDs emitted','30/30 uses','26/26 uses'],
    ['Pilot claims left unverified','41/41','34/34','22/22'],
], [7.2,3.4,3.4,3.4])
placeholder('INSERT FIGURE 2 HERE','One actual synthetic-case report showing claim status, an exact source link and a missing check',1.6)
p('Figure 2. Example synthetic-case report showing claim status, source linkage and unresolved checks. The report illustrates the information presented for review; the findings have not been clinically adjudicated.',style='Caption',align=WD_ALIGN_PARAGRAPH.CENTER)
section(2)
head('5.1 Interpretation of observed results',2)
p('All 18 historical analyses and all 36 pilot analyses completed. Retrieval arms emitted citation identifiers whose identity matched the pinned corpus. This is evidence that source records can be carried through the pipeline and audited. It does not show that each cited passage was relevant to its claim or applicable to the patient. Figure 2 illustrates how these distinctions are presented in the report.')
p('The pilot produced 10/12, 9/12 and 10/12 rank-1 matches. Rank-3 counts were identical. There is no consistent improvement across modes, and the zero graph-candidate count prevents attributing any observed difference to graph evidence. Provider changes further complicate the comparison. The historical gains must also be read in light of reference diagnoses already present in the input text.')
colbreak()
head('5.2 Failure modes and clinical relevance',2)
p('The missed unknown-allergy warning shows that source verification does not correct every upstream failure. The limited authoritative corpus also returned hypertension or amlodipine passages for cases outside its intended domain. An authentic passage can be an unsuitable citation. Keeping claims unverified exposes this distinction, but abstention alone is not proof that the workflow is useful or safe.')
p('All pilot claims remained unverified. The observed contribution is therefore an inspectable review workflow, with conditional proposals and explicit gaps. Broader clinical conclusions require expert-selected cases, controlled models, relevance-labelled retrieval and independent claim review. DECIDE-AI addresses early live evaluation of clinical decision support; the present offline study does not meet that setting.[6]')

# PAGE 6: conclusion, declarations and references.
section(2,True)
head('6. CONCLUSION')
p('This prototype separates exploratory graph and textbook evidence from a narrow citable source corpus, and checks proposed citations independently after generation. The synthetic runs demonstrate report completion, passage-identity traceability and explicit retention of unresolved claims. They also reveal limited graph coverage, a missed missingness warning and provider-confounded comparisons.')
p('The next study should first verify corpus contents and graph relations, then freeze model versions and source snapshots. Relevance-labelled retrieval evaluation and two independent clinical reviewers are needed before assessing claim support or treatment safety. The current evidence supports a feasibility description of the architecture, not diagnostic superiority, autonomous prescribing or improved patient outcomes.')
head('ACKNOWLEDGEMENTS')
p('[If applicable, name contributors who assisted without meeting authorship criteria. Obtain their permission. Otherwise remove this section.]',align=WD_ALIGN_PARAGRAPH.LEFT)
head('DECLARATIONS')
p('Funding: [Insert the actual funding statement and grant identifiers, or confirm that no external funding was received.]',align=WD_ALIGN_PARAGRAPH.LEFT)
p('Conflicts of interest: [Each author must supply a declaration before submission.]',align=WD_ALIGN_PARAGRAPH.LEFT)
p('Ethics and participants: The described evaluations use hand-authored synthetic cases and report no patient or clinician-participant study. No institutional exemption or approval is asserted here. [Add an institutional determination only if documented.]',align=WD_ALIGN_PARAGRAPH.LEFT)
p('Author contributions: [Replace the sample names with actual contributors and specify their roles.]',align=WD_ALIGN_PARAGRAPH.LEFT)
p('Data and code availability: Code, synthetic datasets, source manifest and aggregate snapshots are described in the project repository: https://github.com/Arnavmangarapu/CDSS_SDC2. Generated report files, private review mappings, database volumes and publisher snapshots require separate availability and rights checks.',align=WD_ALIGN_PARAGRAPH.LEFT)
colbreak()
head('REFERENCES')
refs=[
 '1. Lewis P, Perez E, Piktus A, Petroni F, Karpukhin V, Goyal N, et al. Retrieval-augmented generation for knowledge-intensive NLP tasks. Adv Neural Inf Process Syst. 2020;33:9459-9474. Available from: https://arxiv.org/abs/2005.11401',
 '2. Chandak P, Huang K, Zitnik M. Building a knowledge graph to enable precision medicine. Sci Data. 2023;10:67. doi:10.1038/s41597-023-01960-3.',
 '3. pritamdeka. S-PubMedBert-MS-MARCO [Internet]. Hugging Face; [cited 2026 Oct 12]. Available from: https://huggingface.co/pritamdeka/S-PubMedBert-MS-MARCO',
 '4. National Institute for Health and Care Excellence. Hypertension in adults: diagnosis and management. NG136 recommendations [Internet]. London: NICE; 2019 [cited 2026 Oct 12]. Project snapshot updated 2026 Feb 26. Available from: https://www.nice.org.uk/guidance/ng136/chapter/Recommendations',
 '5. National Library of Medicine. Amlodipine besylate tablet [Internet]. DailyMed; 2026 Jan 1 [cited 2026 Oct 12]. Packager: Northwind Health Company, LLC. Set ID: 0c5784b9-5542-144f-e063-6394a90a7ebc. Available from: https://dailymed.nlm.nih.gov/dailymed/lookup.cfm?setid=0c5784b9-5542-144f-e063-6394a90a7ebc',
 '6. Vasey B, Nagendran M, Campbell B, Clifton DA, Collins GS, Denaxas S, et al. Reporting guideline for the early-stage clinical evaluation of decision support systems driven by artificial intelligence: DECIDE-AI. Nat Med. 2022;28:924-933. doi:10.1038/s41591-022-01772-9.',
]
for ref in refs:
    # Permit wraps within URLs without inserting visible hyphens into identifiers.
    display=re.sub(r'https?://\S+',lambda m:m.group().replace('/','/\u200b').replace('?','?\u200b').replace('=','=\u200b'),ref)
    para=p(display,align=WD_ALIGN_PARAGRAPH.LEFT,size=10)
    markdown[-1]=ref+'\n'
    para.paragraph_format.space_after=Pt(9)
    para.paragraph_format.line_spacing=Pt(12)
    para.paragraph_format.keep_together=True
head('SUPPLEMENTARY MATERIAL')
p('If supplied: detailed module diagram; versioned source and dataset manifests; case-level results; retrieval limitations; and a blinded review protocol. Remove this paragraph if no supplementary files accompany the submission. The internal editing page is not supplementary material.',align=WD_ALIGN_PARAGRAPH.LEFT)

OUT.mkdir(parents=True,exist_ok=True)
doc.save(OUT/'ijddt_sample_manuscript.docx')
(OUT/'ijddt_sample_manuscript.md').write_text('\n'.join(markdown),encoding='utf-8')

# PAGE 7: internal revision guide, deliberately excluded from the manuscript copy.
section(1,True)
head('INTERNAL EDITING NOTES')
p('Page 7 - remove before submission',align=WD_ALIGN_PARAGRAPH.LEFT,bold=True)
p('This is a project-specific sample, not an accepted article or a completed submission. The manuscript occupies pages 1-6. Sample authors and institutional names are fictional; author@example.com is a placeholder. No acceptance dates, DOI, journal volume, published page range or ethics decision have been invented.',align=WD_ALIGN_PARAGRAPH.LEFT)
notes=[
 ('Layout and journal requirements','The attached five-page article controls the visual direction: full-width front matter, Times-style text, two-column body and centred figure captions. This sample uses A4, 1.8 cm side margins, 1.7 cm top/bottom margins and 10.5 pt body text. IJDDT online instructions currently request double spacing, at least 2.54 cm margins, a 14 pt bold all-capitals title and superscript numbered references. They place References on a separate sheet. The checklist specifies a 10-page maximum, with exceptions. The user-requested six-page compact layout is a published-style sample; obtain the editor-approved submission format before treating it as compliant. Sources: https://ijddt.com/instructions/ and https://ijddt.com/checklist/ (checked 12 October 2026).'),
 ('Images and page budget','Replace Figure 1 space with a compact architecture diagram at no more than 17.4 cm wide and about 6 cm high; use a landscape arrangement or two panels rather than compressing the earlier portrait image. Preserve the direct pinned-passage-to-verifier arrow. Figure 2 should be an actual synthetic report crop, at most 17.4 cm wide and about 4 cm high. Insert both images inline, replace the slot paragraphs, keep captions attached and recheck all page breaks. Do not add screenshots as extra pages.'),
 ('Technical details to keep stable','The inspected encoder is pritamdeka/S-PubMedBert-MS-MARCO, with 768-dimensional cosine vectors. The textbook recipe still has the 150,000-character cap and 150/20-word chunking. The 0.4 cutoff applies only to textbook retrieval. Do not introduce MedCPT, BioBERT, BM25, RRF, 225,483 chunks or a learned reranker unless a separate verified experiment actually used them. Audit the live collections before asserting actual indexed coverage.'),
 ('Results and permissible conclusions','Keep the historical six cases separate from the 12-case pilot. Pilot rank-1 counts are 10/12, 9/12 and 10/12; graph disease candidates were zero and providers differed for two cases. Citation identity is not citation relevance or clinical support. Do not label these counts patient accuracy or claim safe prescribing. Retrieval Hit@k/Recall@k/MRR and independent clinical adjudication remain unavailable. Update abstract, tables and conclusion together when results change.'),
 ('Where to verify or revise','Use backend/services/vector_retriever.py; ingestion/populate_qdrant.py; backend/services/claim_verifier.py; evaluation/README.md; paper/RESULTS_SNAPSHOT.json; and paper/MULTIDISEASE_PILOT_RESULTS.json. Full runs are under evaluation/runs/ and may be Git ignored. The companion build_sample.py recreates both Word copies; edit its paragraph strings to preserve the page plan. The Markdown companion is useful for text review, but has no fixed pagination.'),
 ('Final author checks and export','Replace five sample authors, affiliations, email and all bracketed declaration fields. Verify authorship, funding, interests, repository accessibility and ethics wording; disclose assisted writing or figure generation if the journal requires it. Check each reference and add focused comparable CDSS studies before claiming novelty. Preserve section breaks and Word styles. The supplied PDFs were exported from these DOCX files with LibreOffice; Times New Roman used its metric-compatible Liberation Serif fallback here. Exact reflow across Microsoft Word, fonts and printer settings cannot be guaranteed. Export again after edits, inspect all pages, and use the supplied PDF as the fixed layout reference.'),
]
for title,text in notes:
    para=p(title+'. '+text,align=WD_ALIGN_PARAGRAPH.LEFT,size=10)
    para.paragraph_format.space_after=Pt(7);para.paragraph_format.line_spacing=Pt(12)
doc.save(OUT/'ijddt_sample_with_editing_notes.docx')
(OUT/'editing_notes.md').write_text('\n'.join(markdown[-9:]),encoding='utf-8')
print('Created manuscript and working DOCX copies in',OUT)
